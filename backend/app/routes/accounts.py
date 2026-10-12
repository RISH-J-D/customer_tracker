import io
from fastapi import APIRouter, Depends, HTTPException, Query, Header, UploadFile, File, Response
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import or_, func
from typing import List, Optional
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from ..database import get_db
from ..models import Account, Client, NotificationItem
from ..schemas import AccountCreate, AccountUpdate, AccountResponse
from ..activity_logger import log_activity

router = APIRouter(prefix="/api/accounts", tags=["Accounts & Projects"])

@router.get("", response_model=List[AccountResponse])
def get_accounts(
    company: Optional[str] = Query(None),
    client_id: Optional[int] = Query(None),
    status: Optional[str] = Query(None),
    search: Optional[str] = Query(None),
    db: Session = Depends(get_db)
):
    query = db.query(Account)

    if company:
        query = query.filter(Account.company.ilike(f"%{company.strip()}%"))
    if client_id is not None:
        query = query.filter(Account.client_id == client_id)
    if status:
        query = query.filter(Account.status == status)
    if search:
        search_fmt = f"%{search.strip()}%"
        query = query.filter(
            or_(
                Account.account_name.ilike(search_fmt),
                Account.company.ilike(search_fmt),
                Account.project_name.ilike(search_fmt),
                Account.client_representative.ilike(search_fmt),
                Account.rep_email.ilike(search_fmt),
                Account.notes.ilike(search_fmt)
            )
        )

    return query.order_by(Account.updated_at.desc()).all()

@router.get("/stats")
def get_account_stats(db: Session = Depends(get_db)):
    """Aggregate statistics for Steel Detailing accounts and active projects"""
    accounts = db.query(Account).all()
    total_accounts = len(accounts)
    total_contract_value = sum(a.deal_value for a in accounts)
    unique_companies = len(set(a.company.lower() for a in accounts if a.company))
    active_projects = sum(1 for a in accounts if a.status in ["Active Project", "In Detailing", "Fabrication"])
    total_tonnage = sum(a.steel_tonnage or 0.0 for a in accounts)

    return {
        "total_accounts": total_accounts,
        "total_contract_value": total_contract_value,
        "unique_companies": unique_companies,
        "active_projects": active_projects,
        "total_tonnage": total_tonnage
    }

@router.get("/export/excel")
def export_accounts_excel(db: Session = Depends(get_db)):
    """Export accounts and projects to a styled Excel workbook"""
    accounts = db.query(Account).order_by(Account.company.asc(), Account.project_name.asc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Accounts & Projects"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    headers = [
        "Account ID", "Company", "Account Name", "Project Name", 
        "Client Representative", "First Name", "Last Name", "Position", 
        "Email", "Phone", "Status", "Contract Value ($)", "Steel Tonnage", 
        "City", "State", "Country", "Notes", "Created Date"
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    status_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")

    for row_idx, a in enumerate(accounts, start=2):
        row_data = [
            a.id,
            a.company or "N/A",
            a.account_name or "N/A",
            a.project_name or "N/A",
            a.client_representative or "N/A",
            a.rep_first_name or "",
            a.rep_last_name or "",
            a.rep_designation or "",
            a.rep_email or "",
            a.rep_phone or "",
            a.status or "Active Project",
            f"${a.deal_value:,.2f}" if a.deal_value else "$0.00",
            f"{a.steel_tonnage:,.1f} Tons" if a.steel_tonnage else "N/A",
            a.city or "",
            a.state or "",
            a.country or "",
            a.notes or "",
            a.created_at.strftime("%Y-%m-%d %H:%M") if a.created_at else ""
        ]
        ws.append(row_data)

        if a.status in ["Active Project", "Accepted", "Completed"]:
            ws.cell(row=row_idx, column=11).fill = status_fill

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if col_idx in [1, 11, 18]:
                cell.alignment = center_align
            elif col_idx in [12, 13]:
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = left_align

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"Accounts_Projects_{datetime.utcnow().strftime('%Y%m%d_%H%M')}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@router.get("/export/csv")
def export_accounts_csv(db: Session = Depends(get_db)):
    """Export accounts and projects to CSV"""
    accounts = db.query(Account).order_by(Account.company.asc(), Account.project_name.asc()).all()
    rows = []
    for a in accounts:
        rows.append({
            "Account ID": a.id,
            "Company": a.company,
            "Account Name": a.account_name,
            "Project Name": a.project_name,
            "Client Representative": a.client_representative,
            "First Name": a.rep_first_name or "",
            "Last Name": a.rep_last_name or "",
            "Position": a.rep_designation or "",
            "Email": a.rep_email or "",
            "Phone": a.rep_phone or "",
            "Status": a.status,
            "Contract Value": a.deal_value,
            "Steel Tonnage": a.steel_tonnage,
            "City": a.city or "",
            "State": a.state or "",
            "Country": a.country or "",
            "Notes": a.notes or "",
            "Created Date": a.created_at.strftime("%Y-%m-%d %H:%M:%S") if a.created_at else ""
        })
    df = pd.DataFrame(rows)
    stream = io.StringIO()
    df.to_csv(stream, index=False)
    filename = f"accounts_projects_{datetime.utcnow().strftime('%Y%m%d')}.csv"
    return Response(
        content=stream.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@router.get("/export/pdf")
def export_accounts_pdf(db: Session = Depends(get_db)):
    """Export accounts and projects summary report to PDF"""
    accounts = db.query(Account).order_by(Account.company.asc(), Account.project_name.asc()).all()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=25,
        leftMargin=25,
        topMargin=25,
        bottomMargin=25
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#0F172A'),
        spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=12
    )
    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#1E293B')
    )
    cell_bold = ParagraphStyle(
        'CellBold',
        parent=styles['Normal'],
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#0F172A'),
        fontName='Helvetica-Bold'
    )

    elements = []
    elements.append(Paragraph("Corporate Accounts & Detailing Projects Report", title_style))
    elements.append(Paragraph(f"Generated on {datetime.utcnow().strftime('%B %d, %Y at %I:%M %p UTC')} | ClientFlow Steel Detailing CRM", subtitle_style))

    total_accounts = len(accounts)
    total_val = sum(a.deal_value for a in accounts)
    total_tons = sum(a.steel_tonnage or 0.0 for a in accounts)
    unique_comps = len(set(a.company.lower() for a in accounts if a.company))

    summary_data = [
        [
            Paragraph("<b>Total Accounts:</b>", cell_style),
            Paragraph(str(total_accounts), cell_bold),
            Paragraph("<b>Companies:</b>", cell_style),
            Paragraph(str(unique_comps), cell_bold),
            Paragraph("<b>Total Contract Value:</b>", cell_style),
            Paragraph(f"${total_val:,.2f}", cell_bold),
            Paragraph("<b>Total Tonnage:</b>", cell_style),
            Paragraph(f"{total_tons:,.0f} Tons", cell_bold)
        ]
    ]
    summary_table = Table(summary_data, colWidths=[90, 50, 70, 50, 110, 90, 80, 80])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#CBD5E1')),
        ('PADDING', (0, 0), (-1, -1), 6),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 12))

    table_headers = [
        Paragraph("<b>Company</b>", cell_bold),
        Paragraph("<b>Project Name</b>", cell_bold),
        Paragraph("<b>Representative</b>", cell_bold),
        Paragraph("<b>Position</b>", cell_bold),
        Paragraph("<b>Contact</b>", cell_bold),
        Paragraph("<b>Status</b>", cell_bold),
        Paragraph("<b>Value</b>", cell_bold),
        Paragraph("<b>Tonnage</b>", cell_bold)
    ]
    table_data = [table_headers]

    for a in accounts:
        contact_str = f"{a.rep_email or ''}<br/>{a.rep_phone or ''}".strip("<br/>")
        table_data.append([
            Paragraph(a.company or "N/A", cell_bold),
            Paragraph(a.project_name or "N/A", cell_style),
            Paragraph(a.client_representative or "N/A", cell_style),
            Paragraph(a.rep_designation or "-", cell_style),
            Paragraph(contact_str or "-", cell_style),
            Paragraph(a.status or "Active", cell_style),
            Paragraph(f"${a.deal_value:,.0f}", cell_style),
            Paragraph(f"{a.steel_tonnage:,.0f} T" if a.steel_tonnage else "-", cell_style)
        ])

    col_widths = [120, 140, 110, 85, 120, 75, 65, 55]
    report_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    report_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
    ]))

    elements.append(report_table)
    doc.build(elements)
    buffer.seek(0)

    filename = f"Accounts_Projects_Report_{datetime.utcnow().strftime('%Y%m%d_%H%M')}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={'Content-Disposition': f'attachment; filename="{filename}"'}
    )

@router.get("/template")
def download_accounts_template():
    """Download an Excel template for importing accounts and projects"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Accounts_Import_Template"

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    center_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    headers = [
        "Account Name", "Company", "Project Name", "Contact Person", 
        "First Name", "Last Name", "Position", "Phone", "Email", 
        "Country", "State", "City", "Deal Value", "Steel Tonnage", 
        "Status", "Notes"
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    sample_rows = [
        [
            "Empire Welding & Construction Inc.", "Empire Welding & Construction Inc.", "Structural Detailing Package", "John Doe",
            "John", "Doe", "Project Manager", "555-019-2831", "j.doe@empirewelding.com",
            "USA", "TX", "Weatherford", 64000, 397,
            "Active Project", "Connection design & Tekla 3D model"
        ],
        [
            "Ace Fabricators", "Ace Fabricators", "Ace Fabricators Detailing Scope", "Jeremy Moore",
            "Jeremy", "Moore", "MainMan", "281-442-0892", "j.moore@acefab.com",
            "USA", "TX", "Houston", 63000, 416,
            "Active Project", "Commercial structural detailing"
        ]
    ]

    for row in sample_rows:
        ws.append(row)

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="Accounts_Import_Template.xlsx"'}
    )

@router.post("/import-excel")
async def import_accounts_excel(
    file: UploadFile = File(...),
    x_user_name: Optional[str] = Header(None),
    x_user_username: Optional[str] = Header(None),
    db: Session = Depends(get_db)
):
    """Import accounts and projects from an Excel or CSV file"""
    filename = file.filename or ""
    if not (filename.endswith(".xlsx") or filename.endswith(".xls") or filename.endswith(".csv")):
        raise HTTPException(status_code=400, detail="Invalid format. Please upload .xlsx, .xls, or .csv")

    contents = await file.read()
    try:
        if filename.endswith(".csv"):
            df = pd.read_csv(io.BytesIO(contents))
        else:
            df = pd.read_excel(io.BytesIO(contents))
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse file: {str(e)}")

    if df.empty:
        raise HTTPException(status_code=400, detail="The uploaded spreadsheet contains no data rows.")

    # Column mapping
    col_map = {}
    for col in df.columns:
        clean = str(col).strip().lower().replace("_", "").replace(" ", "").replace("?", "").replace("(", "").replace(")", "").replace("$", "")
        col_map[clean] = col

    def get_val(row, *aliases):
        for a in aliases:
            clean_a = a.lower().replace("_", "").replace(" ", "").replace("?", "").replace("(", "").replace(")", "").replace("$", "")
            if clean_a in col_map:
                val = row[col_map[clean_a]]
                if pd.notna(val):
                    return str(val).strip()
        return None

    def get_float(row, *aliases):
        val = get_val(row, *aliases)
        if val is None:
            return 0.0
        try:
            cleaned = str(val).replace("$", "").replace(",", "").replace("tons", "").replace("ton", "").strip()
            return float(cleaned)
        except (ValueError, TypeError):
            return 0.0

    imported_count = 0
    skipped_count = 0

    clients_by_name = {c.name.strip().lower(): c.id for c in db.query(Client).all()}
    clients_by_comp = {c.company.strip().lower(): c.id for c in db.query(Client).filter(Client.company.isnot(None)).all()}

    for _, row in df.iterrows():
        account_name = get_val(row, "accountname", "companyname", "company", "account")
        contact_person = get_val(row, "contactperson", "clientrepresentative", "representative", "contact", "name")

        if not account_name and not contact_person:
            skipped_count += 1
            continue

        company = get_val(row, "company", "companyname") or account_name or "Corporate Client"
        account_name = account_name or company
        client_rep = contact_person or "Primary Contact"

        first_name = get_val(row, "firstname", "repfirstname")
        last_name = get_val(row, "lastname", "replastname")
        if not first_name:
            parts = client_rep.strip().split(" ", 1)
            first_name = parts[0]
            last_name = parts[1] if len(parts) > 1 else ""

        designation = get_val(row, "position", "repdesignation", "designation", "title") or None
        phone = get_val(row, "phone", "telephone", "mobile", "tel") or None
        email = get_val(row, "email", "e-mail", "mail") or None
        country = get_val(row, "country", "nation") or None
        state = get_val(row, "state", "province") or None
        city = get_val(row, "city", "town") or None
        notes = get_val(row, "notes", "scope", "comments", "description") or None

        raw_status = get_val(row, "status") or "Active Project"
        status_map = {
            "accepted": "Active Project",
            "active": "Active Project",
            "active project": "Active Project",
            "in detailing": "In Detailing",
            "fabrication": "Fabrication",
            "completed": "Completed",
            "on hold": "On Hold"
        }
        status = status_map.get(raw_status.lower(), raw_status)

        project_name = get_val(row, "projectname", "project", "order", "jobname")
        if not project_name:
            project_name = f"{company} Detailing Project"

        deal_value = get_float(row, "dealvalue", "contractvalue", "value", "price", "estvalue")
        tonnage = get_float(row, "steeltonnage", "tonnage", "tons") or None

        client_id = clients_by_comp.get(company.strip().lower()) or clients_by_name.get(client_rep.strip().lower())

        account = Account(
            client_id=client_id,
            account_name=account_name,
            company=company,
            project_name=project_name,
            client_representative=client_rep,
            rep_first_name=first_name,
            rep_last_name=last_name,
            rep_email=email,
            rep_phone=phone,
            rep_designation=designation,
            deal_value=deal_value,
            steel_tonnage=tonnage,
            status=status,
            city=city,
            state=state,
            country=country,
            notes=notes
        )
        db.add(account)
        imported_count += 1

    db.commit()

    log_activity(
        db=db,
        action="IMPORT_EXCEL",
        entity_type="Account",
        entity_name="Accounts Import",
        description=f"Imported {imported_count} accounts/projects from '{filename}' (Skipped {skipped_count})",
        user_name=x_user_name,
        user_username=x_user_username
    )

    return {
        "status": "success",
        "imported": imported_count,
        "skipped": skipped_count,
        "message": f"Successfully imported {imported_count} accounts and projects."
    }

@router.post("", response_model=AccountResponse)
def create_account(
    account_in: AccountCreate,
    x_user_name: Optional[str] = Header(None),
    x_user_username: Optional[str] = Header(None),
    db: Session = Depends(get_db)
):
    data = account_in.model_dump()
    # Auto-synchronize first/last name with full client_representative string
    first = (data.get("rep_first_name") or "").strip()
    last = (data.get("rep_last_name") or "").strip()
    if first or last:
        data["client_representative"] = f"{first} {last}".strip()
    elif data.get("client_representative"):
        parts = data["client_representative"].strip().split(" ", 1)
        data["rep_first_name"] = parts[0]
        data["rep_last_name"] = parts[1] if len(parts) > 1 else ""

    account = Account(**data)
    db.add(account)
    db.commit()
    db.refresh(account)

    # Activity log
    is_from_client = account.client_id is not None
    action_type = "CONVERT_LEAD_TO_ACCOUNT" if is_from_client else "CREATE_ACCOUNT"
    description = (
        f"Created Account '{account.account_name}' under company '{account.company}' "
        f"for project '{account.project_name}' (Rep: {account.client_representative}, Value: ${account.deal_value:,.2f})"
    )

    log_activity(
        db=db,
        action=action_type,
        entity_type="Account",
        entity_id=account.id,
        entity_name=f"{account.account_name} - {account.project_name}",
        description=description,
        changes={
            "company": account.company,
            "project_name": account.project_name,
            "client_representative": account.client_representative,
            "deal_value": account.deal_value,
            "status": account.status,
            "steel_tonnage": account.steel_tonnage
        },
        user_name=x_user_name,
        user_username=x_user_username
    )

    # Notify
    notif = NotificationItem(
        title=f"🏢 New Account / Project: {account.project_name}",
        message=f"Account '{account.account_name}' created for {account.company} with Rep {account.client_representative}.",
        type="account_created",
        client_id=account.client_id
    )
    db.add(notif)
    db.commit()

    return account

@router.get("/{account_id}", response_model=AccountResponse)
def get_account(account_id: int, db: Session = Depends(get_db)):
    account = db.query(Account).filter(Account.id == account_id).first()
    if not account:
        raise HTTPException(status_code=404, detail="Account not found")
    return account

@router.put("/{account_id}", response_model=AccountResponse)
def update_account(
    account_id: int,
    account_update: AccountUpdate,
    x_user_name: Optional[str] = Header(None),
    x_user_username: Optional[str] = Header(None),
    db: Session = Depends(get_db)
):
    account = db.query(Account).filter(Account.id == account_id).first()
    if not account:
        raise HTTPException(status_code=404, detail="Account not found")

    update_data = account_update.model_dump(exclude_unset=True)
    if "rep_first_name" in update_data or "rep_last_name" in update_data:
        first = (update_data.get("rep_first_name", account.rep_first_name) or "").strip()
        last = (update_data.get("rep_last_name", account.rep_last_name) or "").strip()
        update_data["client_representative"] = f"{first} {last}".strip()
    elif "client_representative" in update_data and update_data["client_representative"]:
        parts = update_data["client_representative"].strip().split(" ", 1)
        update_data["rep_first_name"] = parts[0]
        update_data["rep_last_name"] = parts[1] if len(parts) > 1 else ""

    diffs = {}
    for field, new_val in update_data.items():
        old_val = getattr(account, field, None)
        if str(old_val) != str(new_val):
            diffs[field] = {"old": old_val, "new": new_val}

    for field, val in update_data.items():
        setattr(account, field, val)

    account.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(account)

    if diffs:
        log_activity(
            db=db,
            action="UPDATE_ACCOUNT",
            entity_type="Account",
            entity_id=account.id,
            entity_name=f"{account.account_name} ({account.company})",
            description=f"Updated Account '{account.account_name}': {list(diffs.keys())}",
            changes=diffs,
            user_name=x_user_name,
            user_username=x_user_username
        )

    return account

@router.delete("/{account_id}")
def delete_account(
    account_id: int,
    x_user_name: Optional[str] = Header(None),
    x_user_username: Optional[str] = Header(None),
    db: Session = Depends(get_db)
):
    account = db.query(Account).filter(Account.id == account_id).first()
    if not account:
        raise HTTPException(status_code=404, detail="Account not found")

    acc_name = account.account_name
    proj_name = account.project_name
    company = account.company

    db.delete(account)
    db.commit()

    log_activity(
        db=db,
        action="DELETE_ACCOUNT",
        entity_type="Account",
        entity_id=account_id,
        entity_name=f"{acc_name} ({company})",
        description=f"Deleted Account '{acc_name}' - Project '{proj_name}' ({company})",
        user_name=x_user_name,
        user_username=x_user_username
    )

    return {"message": "Account deleted successfully", "id": account_id}
